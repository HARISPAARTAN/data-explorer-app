"""core/io.py: detection, loading, header cleanup, missing values, dates, errors and limits."""

from __future__ import annotations

import dataclasses
import io
import json
import re
import struct
import sys
import time
import zipfile
from collections.abc import Callable
from typing import Any

import pandas as pd
import pytest
from openpyxl import Workbook

from data_explorer.core import io as dio
from data_explorer.core.io import (
    BROWSER_LIMITS,
    SAMPLES,
    SERVER_LIMITS,
    LoadError,
    LoadOptions,
    MiB,
    current_limits,
    dataset_id,
    detect_decimal,
    detect_delimiter,
    detect_encoding,
    detect_header_row,
    detect_sheet,
    list_sheets,
    load_sample,
    load_table,
    sample_bytes,
)

Read = Callable[[str], bytes]


def make_xlsx(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def patch_member(data: bytes, member: str, change: Callable[[bytes], bytes]) -> bytes:
    """The zip with one member rewritten, as the same container otherwise."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            content = source.read(item.filename)
            target.writestr(item, change(content) if item.filename == member else content)
    return out.getvalue()


def damage_member(data: bytes, member: str) -> bytes:
    """Flip bytes inside the compressed stream of one member, leaving the zip directory intact."""
    info = zipfile.ZipFile(io.BytesIO(data)).getinfo(member)
    name_length, extra_length = struct.unpack("<HH", data[info.header_offset + 26 : info.header_offset + 30])
    start = info.header_offset + 30 + name_length + extra_length
    damaged = bytearray(data)
    for position in range(start + 4, start + 12):
        damaged[position] ^= 0xFF
    return bytes(damaged)


SHEET = "xl/worksheets/sheet1.xml"
DIMENSION = re.compile(rb'<dimension ref="[^"]*"\s*/>')


@pytest.mark.parametrize(
    ("name", "encoding", "delimiter", "decimal"),
    [
        ("utf8_comma.csv", "utf-8", ",", "."),
        ("utf8_bom_semicolon_decimal_comma.csv", "utf-8-sig", ";", ","),
        ("cp1252_semicolon.csv", "cp1252", ";", "."),
        ("latin1_tab.tsv", "latin-1", "\t", "."),
        ("pipe_delimited.txt", "utf-8", "|", "."),
    ],
)
def test_detects_encoding_delimiter_and_decimal(
    fixture_bytes: Read, name: str, encoding: str, delimiter: str, decimal: str
) -> None:
    df, report = load_table(fixture_bytes(name), name)
    assert (report.encoding, report.delimiter, report.decimal) == (encoding, delimiter, decimal)
    assert report.file_type == "csv" and report.header_row == 0
    assert (report.n_rows, report.n_cols) == df.shape
    assert report.sheet is None and report.sheets == ()


def test_values_are_parsed_per_dialect(fixture_bytes: Read) -> None:
    df, _ = load_table(fixture_bytes("utf8_bom_semicolon_decimal_comma.csv"), "a.csv")
    assert df["amount"].dtype == "float64" and df["amount"].iloc[0] == pytest.approx(1234.56)
    assert df["day"].dtype == "datetime64[us]" and df["day"].iloc[0] == pd.Timestamp("2025-03-01")
    assert df["units"].dtype == "int64"
    df, _ = load_table(fixture_bytes("utf8_comma.csv"), "b.csv")
    assert df["label"].iloc[0] == "Caffè" and df["amount"].iloc[1] == pytest.approx(11.5)
    df, _ = load_table(fixture_bytes("cp1252_semicolon.csv"), "c.csv")
    assert df["note"].iloc[0] == "Prezzo in €" and df["price_eur"].dtype == "float64"
    df, _ = load_table(fixture_bytes("latin1_tab.tsv"), "d.tsv")
    assert df["name"].iloc[0] == "Kühl" and df["note"].iloc[0] == "ctrl\x81"


def test_xlsx_picks_the_larger_sheet_and_header_row_four(fixture_bytes: Read) -> None:
    data = fixture_bytes("header_offset_two_sheets.xlsx")
    assert list_sheets(data) == ("Summary", "Data")
    assert detect_sheet(data) == "Data"
    df, report = load_table(data, "offset.xlsx")
    assert (report.file_type, report.sheet, report.sheets, report.header_row) == (
        "xlsx",
        "Data",
        ("Summary", "Data"),
        3,
    )
    assert list(df.columns) == ["month", "region", "amount", "units"]
    assert len(df) == 30 and df["units"].dtype == "int64" and df["amount"].dtype == "float64"
    assert report.encoding is None and report.delimiter is None and report.decimal is None


def test_xlsx_sheet_option_and_missing_sheet(fixture_bytes: Read) -> None:
    data = fixture_bytes("header_offset_two_sheets.xlsx")
    df, report = load_table(data, "offset.xlsx", LoadOptions(sheet="Summary"))
    assert report.sheet == "Summary" and list(df.columns) == ["Summary", "column_2"] and df.shape == (1, 2)
    with pytest.raises(LoadError) as caught:
        load_table(data, "offset.xlsx", LoadOptions(sheet="Nope"))
    assert caught.value.code == "sheet_not_found" and "Summary, Data" in caught.value.message


def test_renames_blank_and_repeated_headers(fixture_bytes: Read) -> None:
    df, report = load_table(fixture_bytes("duplicate_and_blank_headers.csv"), "h.csv")
    assert list(df.columns) == ["id", "name", "column_3", "name_2", "amount", "name_3"]
    assert report.renamed_columns == (
        ("", "column_3"),
        ("name", "name_2"),
        (" amount ", "amount"),
        ("name", "name_3"),
    )
    assert "Column 7 has no header and no values and was removed." in report.warnings
    assert df["amount"].tolist() == [10, 20, 30, 40, 50]


def test_whitespace_cells_become_missing_and_text_is_kept(fixture_bytes: Read) -> None:
    df, report = load_table(fixture_bytes("whitespace_cells.csv"), "w.csv")
    assert report.blank_to_missing == 4
    assert df["name"].isna().sum() == 1
    assert df["score"].dtype == "Int64" and df["score"].isna().sum() == 1
    assert df["note"].isna().sum() == 3 and df["note"].iloc[0] == "  padded  "


def test_date_formats_and_no_month_first(fixture_bytes: Read) -> None:
    df, report = load_table(fixture_bytes("dates_mixed.csv"), "dates.csv")
    parsed = ("iso", "day_first", "dotted", "iso_minutes", "iso_t", "offset")
    for column in parsed:
        assert df[column].dtype == "datetime64[us]", column
    assert set(report.date_columns) == set(parsed)
    assert pd.api.types.is_string_dtype(df["month_first"]) and pd.api.types.is_string_dtype(df["junk"])
    assert df["day_first"].isna().sum() == 1 and df["day_first"].iloc[0] == pd.Timestamp("2025-03-10")
    assert df["dotted"].iloc[0] == pd.Timestamp("2025-03-10")
    assert df["iso_minutes"].iloc[0] == pd.Timestamp("2025-03-01 08:15")
    assert df["iso_t"].iloc[0] == pd.Timestamp("2025-03-01 08:15:30")
    assert df["offset"].iloc[0] == pd.Timestamp("2025-03-01 07:00")
    assert df["junk"].isna().sum() == 3
    assert any("'day_first'" in w and "1 of 20 values" in w for w in report.warnings)
    assert any("'offset'" in w and "UTC" in w for w in report.warnings)


def test_ambiguous_dates_parse_day_first_and_month_first_stays_text() -> None:
    df, _ = load_table(b"d\n01/02/2025\n03/04/2025\n", "a.csv")
    assert df["d"].iloc[0] == pd.Timestamp("2025-02-01")
    df, report = load_table(b"d\n12/31/2024\n01/15/2025\n", "b.csv")
    assert pd.api.types.is_string_dtype(df["d"]) and report.date_columns == ()


@pytest.mark.parametrize(
    ("first", "second", "expected"),
    [
        ("2025-03-01 08:15:30", "2025-03-02 09:45:10", "2025-03-01 08:15:30"),
        ("2025-03-01 08:15:30.25", "2025-03-02 09:45:10.5", "2025-03-01 08:15:30.25"),
        ("2025-03-01T08:15", "2025-03-02T09:45", "2025-03-01 08:15:00"),
        ("2025-03-01T08:15:30.125", "2025-03-02T09:45:10.5", "2025-03-01 08:15:30.125"),
        ("10.03.2025 08:15", "11.03.2025 09:45", "2025-03-10 08:15:00"),
        ("10.03.2025 08:15:30", "11.03.2025 09:45:10", "2025-03-10 08:15:30"),
        ("10/03/2025 08:15:30", "11/03/2025 09:45:10", "2025-03-10 08:15:30"),
    ],
)
def test_common_timestamp_formats_become_datetime_columns(first: str, second: str, expected: str) -> None:
    df, report = load_table(f"t\n{first}\n{second}\n".encode(), "t.csv")
    assert df["t"].dtype == "datetime64[us]" and df["t"].iloc[0] == pd.Timestamp(expected)
    assert df["t"].notna().all() and report.date_columns == ("t",) and report.warnings == ()


def test_dates_in_an_unknown_format_stay_text_with_a_warning() -> None:
    cases = (
        b"d\n1-3-2025\n2-3-2025\n",
        b"d\n12/31/2024\n01/15/2025\n",
        b"d\n2025-03-01\n2025-03-01 08:15:30\n",
    )
    for data in cases:
        df, report = load_table(data, "d.csv")
        assert pd.api.types.is_string_dtype(df["d"]) and report.date_columns == ()
        assert report.warnings == ("Column 'd' looks like dates but no known format matched; kept as text.",)
    df, report = load_table(b"firmware,note\n2.4.1,soon\n2.4.1,later\n", "f.csv")
    assert report.warnings == () and pd.api.types.is_string_dtype(df["firmware"])


def test_booleans_leading_zeros_and_long_integers() -> None:
    df, _ = load_table(b"flag,code,big,n\nTRUE,007,1234567890123456,1\nfalse,010,2,\n", "t.csv")
    assert df["flag"].dtype == "bool" and df["flag"].tolist() == [True, False]
    assert df["code"].tolist() == ["007", "010"]
    assert pd.api.types.is_string_dtype(df["big"])
    assert df["n"].dtype == "Int64"


def test_decimal_comma_with_thousands_and_negatives() -> None:
    df, report = load_table(b"a;b\n-1.234,5;7\n2,25;-3\n", "n.csv")
    assert report.decimal == ","
    assert df["a"].tolist() == [-1234.5, 2.25] and df["b"].tolist() == [7, -3]


def test_decimal_comma_does_not_read_a_zero_point_as_thousands() -> None:
    options = LoadOptions(delimiter=";", decimal=",")
    df, report = load_table(b"x;y\n0.125;0.250\n0.5;0.75\n", "z.csv", options)
    assert report.decimal == ","
    assert df["x"].tolist() == ["0.125", "0.5"] and df["y"].tolist() == ["0.250", "0.75"]
    assert pd.api.types.is_string_dtype(df["x"]) and pd.api.types.is_string_dtype(df["y"])
    df, _ = load_table(b"x;y\n-0.125,5;0.125,5\n", "z.csv", options)
    assert df["x"].tolist() == ["-0.125,5"] and df["y"].tolist() == ["0.125,5"]
    df, _ = load_table(
        b"i;f;n\n1.250;1.125,5;0\n12.500;123.456.789,25;-7\n-1.000;-1.000,5;100\n", "t.csv", options
    )
    assert df["i"].tolist() == [1250, 12500, -1000] and df["i"].dtype == "int64"
    assert df["f"].tolist() == [1125.5, 123456789.25, -1000.5]
    assert df["n"].tolist() == [0, -7, 100]


@pytest.mark.parametrize(
    ("fixture", "file_name", "code"),
    [
        ("empty.csv", "empty.csv", "empty_file"),
        ("legacy.xls", "legacy.xls", "unsupported_type"),
        ("not_a_table.bin", "not_a_table.bin", "unsupported_type"),
        ("not_a_table.bin", "not_a_table.csv", "decode_failed"),
        ("not_a_table.bin", "not_a_table.xlsx", "parse_failed"),
    ],
)
def test_error_codes_for_bad_files(fixture_bytes: Read, fixture: str, file_name: str, code: str) -> None:
    with pytest.raises(LoadError) as caught:
        load_table(fixture_bytes(fixture), file_name)
    assert caught.value.code == code
    assert file_name in caught.value.message and str(caught.value) == caught.value.message


def test_blank_text_and_ragged_rows() -> None:
    with pytest.raises(LoadError) as caught:
        load_table(b" \n\n", "blank.csv")
    assert caught.value.code == "no_columns"
    with pytest.raises(LoadError) as caught:
        load_table(b"a,b\n1,2\n1,2,3\n", "ragged.csv")
    assert caught.value.code == "parse_failed"


def test_limits_raise_with_the_limit_in_the_message(tight_limits: Any) -> None:
    cases = [
        (b"a,b\n" + b"1,2\n" * 600, "too_large", "2,000 bytes"),
        (b"a\n" + b"1\n" * 60, "too_many_rows", "the limit is 50 rows"),
        (b"a,b,c,d,e,f\n1,2,3,4,5,6\n", "too_many_columns", "the limit is 5 columns"),
    ]
    for data, code, text in cases:
        with pytest.raises(LoadError) as caught:
            load_table(data, "big.csv", limits=tight_limits)
        assert caught.value.code == code and text in caught.value.message


@pytest.fixture
def parsed_rows(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """The row count of every finished ``pd.read_csv`` call; a call that fails is not recorded."""
    rows: list[int] = []
    real = pd.read_csv

    def spy(*args: Any, **kwargs: Any) -> pd.DataFrame:
        frame = real(*args, **kwargs)
        rows.append(len(frame))
        return frame

    monkeypatch.setattr(pd, "read_csv", spy)
    return rows


def test_csv_with_too_many_columns_is_refused_before_the_parse(
    tight_limits: Any, parsed_rows: list[int]
) -> None:
    wide = dataclasses.replace(tight_limits, max_bytes=MiB)
    data = ",".join(["a"] * 300_000).encode() + b"\n1\n"
    with pytest.raises(LoadError) as caught:
        load_table(data, "wide.csv", limits=wide)
    assert caught.value.code == "too_many_columns"
    assert caught.value.message == "wide.csv has 300,000 columns; the limit is 5 columns."
    assert parsed_rows == []
    assert load_table(b"a,b,c,d,e\n1,2,3,4,5\n", "five.csv", limits=wide)[0].shape == (1, 5)
    assert parsed_rows == [2]


def test_a_header_cut_by_the_detection_sample_is_still_refused_before_the_parse(
    tight_limits: Any, parsed_rows: list[int]
) -> None:
    wide = dataclasses.replace(tight_limits, max_bytes=4 * MiB)
    data = ",".join(["a"] * 700_000).encode() + b"\n"  # 1.4 MB, longer than the first MiB that is read
    with pytest.raises(LoadError) as caught:
        load_table(data, "wide.csv", limits=wide)
    assert caught.value.code == "too_many_columns" and "has at least " in caught.value.message
    assert caught.value.message.endswith("columns; the limit is 5 columns.")
    assert parsed_rows == []


def test_csv_with_too_many_rows_stops_one_row_past_the_limit(
    tight_limits: Any, parsed_rows: list[int]
) -> None:
    ok = b"a\n" + b"1\n" * 50
    assert len(load_table(ok, "ok.csv", limits=tight_limits)[0]) == 50
    assert parsed_rows == [51]  # the header and 50 rows
    parsed_rows.clear()
    data = b"a\n" + b"1\n" * 900  # 1,800 bytes: the whole file would be 900 rows
    with pytest.raises(LoadError) as caught:
        load_table(data, "long.csv", limits=tight_limits)
    assert caught.value.code == "too_many_rows"
    assert caught.value.message == "long.csv has more than 50 rows; the limit is 50 rows."
    assert parsed_rows == [52]  # the header, 50 rows and the one that shows the limit is passed
    with pytest.raises(LoadError, match="more than 50 rows"):
        load_table(b"a\n" + b"1\n" * 51, "one_too_many.csv", limits=tight_limits)


def test_a_header_after_title_rows_counts_its_own_fields(tight_limits: Any) -> None:
    wide = dataclasses.replace(tight_limits, max_bytes=MiB)
    title = b"Report,,,,,,,,,,\n"  # 11 fields in a row that is not the header
    assert load_table(title + b"a,b\n1,2\n", "t.csv", LoadOptions(header_row=1), wide)[0].shape == (1, 2)


def test_a_byte_order_mark_left_in_the_text_is_dropped() -> None:
    df, report = load_table(b"\xef\xbb\xbfa,b\n1,2\n", "bom.csv", LoadOptions(encoding="utf-8"))
    assert list(df.columns) == ["a", "b"] and report.renamed_columns == ()


def test_inflated_xlsx_dimension_does_not_inflate_the_read() -> None:
    rows: list[list[Any]] = [["id", "ratio", "label"]] + [[i, i / 4, f"x{i}"] for i in range(2000)]
    honest = make_xlsx(rows)
    inflated = patch_member(honest, SHEET, lambda xml: DIMENSION.sub(b'<dimension ref="A1:XFD2001"/>', xml))
    assert b'ref="A1:XFD2001"' in zipfile.ZipFile(io.BytesIO(inflated)).read(SHEET)
    expected, _ = load_table(honest, "honest.xlsx")
    started = time.perf_counter()
    df, report = load_table(inflated, "inflated.xlsx")
    assert time.perf_counter() - started < 1.0
    assert df.shape == (2000, 3) and (report.n_rows, report.n_cols) == (2000, 3) and report.warnings == ()
    pd.testing.assert_frame_equal(df, expected)
    assert detect_sheet(inflated) == "Sheet"


def test_xlsx_row_limit_counts_filled_rows_and_stops_reading(tight_limits: Any) -> None:
    tight_limits = dataclasses.replace(tight_limits, max_bytes=MiB)  # a workbook is bigger than 2,000 bytes
    header: list[Any] = ["a", "b"]
    df, _ = load_table(make_xlsx([header] + [[i, i] for i in range(50)]), "fits.xlsx", limits=tight_limits)
    assert len(df) == 50
    with pytest.raises(LoadError) as caught:
        load_table(make_xlsx([header] + [[i, i] for i in range(51)]), "long.xlsx", limits=tight_limits)
    assert caught.value.code == "too_many_rows" and "the limit is 50 rows" in caught.value.message
    blanks: list[list[Any]] = [header] + [[i, i] for i in range(40)] + [["  ", None]] * 30
    df, report = load_table(make_xlsx(blanks), "blanks.xlsx", limits=tight_limits)
    assert len(df) == 40 and report.n_rows == 40


def test_xlsx_column_limit_uses_the_cells_that_are_filled(tight_limits: Any) -> None:
    tight_limits = dataclasses.replace(tight_limits, max_bytes=MiB)  # a workbook is bigger than 2,000 bytes
    for cells in (list("abcdef"), ["a", None, None, None, None, None, None, None, None, "j"]):
        with pytest.raises(LoadError) as caught:
            load_table(make_xlsx([cells, [1] * len(cells)]), "wide.xlsx", limits=tight_limits)
        assert caught.value.code == "too_many_columns" and "the limit is 5 columns" in caught.value.message
    inflated = patch_member(
        make_xlsx([list("abcde"), [1, 2, 3, 4, 5]]),
        SHEET,
        lambda xml: DIMENSION.sub(b'<dimension ref="A1:XFD2"/>', xml),
    )
    df, _ = load_table(inflated, "five.xlsx", limits=tight_limits)
    assert df.shape == (1, 5)


SHARED_STRINGS = "xl/sharedStrings.xml"


def with_part(data: bytes, name: str, body: bytes) -> bytes:
    """The workbook plus one more member, deflated."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr(name, body)
    return out.getvalue()


def with_shared_strings(data: bytes, megabytes: int) -> bytes:
    """The workbook plus a shared-strings part that unpacks to ``megabytes`` MiB and deflates to a few KiB."""
    entry = b"<si><t>x</t></si>"
    return with_part(data, SHARED_STRINGS, b"<sst>" + entry * (megabytes * MiB // len(entry)) + b"</sst>")


def test_workbook_declaring_a_huge_part_is_refused_before_openpyxl_reads_it(
    tight_limits: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    limits = dataclasses.replace(tight_limits, max_bytes=MiB)  # 10 MB once unpacked
    bomb = with_shared_strings(make_xlsx([["a", "b"], [1, 2]]), 12)
    assert len(bomb) < MiB // 4  # a few KiB of Deflate stand for 12 MB

    def never(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("openpyxl was asked to open the workbook")

    monkeypatch.setattr(dio, "load_workbook", never)
    calls = {
        "load_table": lambda: load_table(bomb, "bomb.xlsx", limits=limits),
        "list_sheets": lambda: list_sheets(bomb, "bomb.xlsx", limits),
        "detect_sheet": lambda: detect_sheet(bomb, limits),
    }
    for label, call in calls.items():
        started = time.perf_counter()
        with pytest.raises(LoadError) as caught:
            call()
        assert caught.value.code == "too_large", label
        assert time.perf_counter() - started < 5.0, label  # generous: the work is reading a zip directory
    with pytest.raises(LoadError) as caught:
        load_table(bomb, "bomb.xlsx", limits=limits)
    assert caught.value.message == (
        "bomb.xlsx unpacks to 12.0 MB; the limit is 10 MB unpacked (10 times the upload limit)."
    )
    assert str(caught.value) == caught.value.message


def test_unpacked_size_cap_is_ten_times_the_upload_limit(tight_limits: Any) -> None:
    data = make_xlsx([["a", "b"], [1, 2]])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        declared = sum(item.file_size for item in archive.infolist())
    enough = dataclasses.replace(
        tight_limits, max_bytes=-(-declared // 10)
    )  # the smallest limit that allows it
    assert enough.max_bytes * 10 >= declared
    assert list_sheets(data, "fits.xlsx", enough) == ("Sheet",)
    assert detect_sheet(data, enough) == "Sheet"
    short = dataclasses.replace(enough, max_bytes=enough.max_bytes - 1)
    assert short.max_bytes * 10 < declared
    with pytest.raises(LoadError) as caught:
        list_sheets(data, "short.xlsx", short)
    assert caught.value.code == "too_large" and "short.xlsx unpacks to" in caught.value.message


def shared_strings_of_size(size: int) -> bytes:
    """An ``sst`` element padded with spaces to exactly ``size`` bytes; it deflates to a few hundred bytes."""
    return b"<sst>" + b" " * (size - len(b"<sst></sst>")) + b"</sst>"


@pytest.mark.parametrize(
    ("limits", "megabytes", "limit"),
    [(SERVER_LIMITS, 60, "50 MB"), (BROWSER_LIMITS, 25, "20 MB")],
    ids=["server", "browser"],
)
def test_a_huge_shared_strings_part_is_refused_at_once(
    monkeypatch: pytest.MonkeyPatch, limits: dio.Limits, megabytes: int, limit: str
) -> None:
    declared = megabytes * MiB
    assert declared <= dio.XLSX_EXPANSION * limits.max_bytes  # only the per-part cap can refuse it
    bomb = with_part(make_xlsx([["a", "b"], [1, 2]]), SHARED_STRINGS, shared_strings_of_size(declared))
    assert len(bomb) < MiB // 4  # a few KiB of Deflate stand for the whole part

    def never(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("openpyxl was asked to open the workbook")

    monkeypatch.setattr(dio, "load_workbook", never)
    monkeypatch.setattr(dio, "current_limits", lambda: limits)  # what a call without limits uses
    calls = {  # label: (name in the message, call)
        "load_table": ("bomb.xlsx", lambda: load_table(bomb, "bomb.xlsx", limits=limits)),
        "list_sheets": ("bomb.xlsx", lambda: list_sheets(bomb, "bomb.xlsx")),
        "detect_sheet": ("the workbook", lambda: detect_sheet(bomb)),
    }
    for label, (name, call) in calls.items():
        started = time.perf_counter()
        with pytest.raises(LoadError) as caught:
            call()
        assert time.perf_counter() - started < 2.0, label
        assert caught.value.code == "too_large", label
        assert caught.value.message == (
            f"{name} has xl/sharedStrings.xml unpacking to {megabytes} MB; "
            f"the limit is {limit} unpacked for that part (the upload limit)."
        ), label


def test_shared_strings_part_may_declare_exactly_the_upload_limit(tight_limits: Any) -> None:
    limits = dataclasses.replace(tight_limits, max_bytes=100_000)  # the total cap, 1,000,000, is far away
    base = make_xlsx([["a", "b"], [1, 2]])
    at_limit = with_part(base, SHARED_STRINGS, shared_strings_of_size(100_000))
    df, _ = load_table(at_limit, "limit.xlsx", limits=limits)
    assert list(df.columns) == ["a", "b"] and df.shape == (1, 2)
    over = with_part(base, SHARED_STRINGS, shared_strings_of_size(100_001))
    for call in (
        lambda: load_table(over, "over.xlsx", limits=limits),
        lambda: list_sheets(over, "over.xlsx", limits),
        lambda: detect_sheet(over, limits),
    ):
        with pytest.raises(LoadError) as caught:
            call()
        assert caught.value.code == "too_large"
        assert "unpacking to 100,001 bytes; the limit is 100,000 bytes" in caught.value.message


def test_a_normal_workbook_still_loads_under_the_default_limits(fixture_bytes: Read) -> None:
    data = fixture_bytes("header_offset_two_sheets.xlsx")
    assert list_sheets(data) == ("Summary", "Data")
    assert load_table(data, "offset.xlsx")[1].n_rows == 30
    with pytest.raises(LoadError) as caught:
        list_sheets(fixture_bytes("not_a_table.bin"), "not_a_table.xlsx")
    assert caught.value.code == "parse_failed"


def test_damaged_xlsx_parts_are_load_errors() -> None:
    good = make_xlsx([["a", "b"], [1, 2], [3, 4]])
    cases = {
        "garbage sheet": patch_member(good, SHEET, lambda _: b"<<< not xml"),
        "truncated sheet": patch_member(good, SHEET, lambda xml: xml[: xml.index(b'<row r="3"') + 20]),
        "damaged stream": damage_member(good, SHEET),
        "garbage workbook part": patch_member(good, "xl/workbook.xml", lambda _: b"\x00\x01 nonsense"),
    }
    for label, data in cases.items():
        with pytest.raises(LoadError) as caught:
            load_table(data, "damaged.xlsx")
        assert caught.value.code == "parse_failed" and "damaged.xlsx" in caught.value.message, label
        with pytest.raises(LoadError) as caught:
            detect_sheet(data)
        assert caught.value.code == "parse_failed", label


def test_xlsx_without_worksheets_has_no_columns() -> None:
    empty = patch_member(
        make_xlsx([["a"], [1]]),
        "xl/workbook.xml",
        lambda xml: re.sub(rb"<sheets>.*</sheets>", b"<sheets/>", xml),
    )
    assert list_sheets(empty) == ()
    for call in (
        lambda: load_table(empty, "none.xlsx"),
        lambda: load_table(empty, "none.xlsx", LoadOptions(sheet="Sheet")),
        lambda: detect_sheet(empty),
    ):
        with pytest.raises(LoadError) as caught:
            call()
        assert caught.value.code == "no_columns" and "no worksheets" in caught.value.message


def test_check_order_empty_then_size_then_type(tight_limits: Any) -> None:
    with pytest.raises(LoadError, match="is empty") as caught:
        load_table(b"", "x.xls")
    assert caught.value.code == "empty_file"
    with pytest.raises(LoadError) as caught:
        load_table(b"x" * 3_000, "x.xls", limits=tight_limits)
    assert caught.value.code == "too_large"


def test_header_only_gives_zero_rows(fixture_bytes: Read) -> None:
    df, report = load_table(fixture_bytes("header_only.csv"), "header_only.csv")
    assert df.shape == (0, 3) and report.n_rows == 0 and list(df.columns) == ["a", "b", "c"]


def test_options_override_detection(fixture_bytes: Read) -> None:
    options = LoadOptions(encoding="utf-8-sig", delimiter=";", decimal=",", header_row=0)
    df, report = load_table(fixture_bytes("utf8_bom_semicolon_decimal_comma.csv"), "a.csv", options)
    assert report.delimiter == ";" and df["amount"].iloc[0] == pytest.approx(1234.56)
    with pytest.raises(LoadError) as caught:
        load_table(fixture_bytes("cp1252_semicolon.csv"), "c.csv", LoadOptions(encoding="utf-8"))
    assert caught.value.code == "decode_failed" and "Reading options" in caught.value.message


def test_load_options_validate_and_round_trip() -> None:
    assert (
        LoadOptions().to_json()
        == '{"decimal":null,"delimiter":null,"encoding":null,"header_row":null,"sheet":null}'
    )
    options = LoadOptions(encoding="cp1252", delimiter="\t", decimal=",", sheet="Data", header_row=3)
    assert LoadOptions.from_json(options.to_json()) == options
    for bad in (
        {"delimiter": ":"},
        {"encoding": "utf-16"},
        {"decimal": ";"},
        {"header_row": 20},
        {"header_row": -1},
    ):
        with pytest.raises(ValueError):
            LoadOptions(**bad)


def test_detect_encoding() -> None:
    assert detect_encoding(b"\xef\xbb\xbfa,b") == "utf-8-sig"
    assert detect_encoding("é,ü".encode()) == "utf-8"
    assert detect_encoding(b"a" * (MiB - 1) + "é".encode()) == "utf-8"
    assert detect_encoding("€".encode("cp1252")) == "cp1252"
    assert detect_encoding(b"\x81abc") == "latin-1"


def test_detect_delimiter_and_decimal() -> None:
    assert detect_delimiter("a;b;c\n1;2;3\n") == ";"
    assert detect_delimiter("a|b\n1|2\n") == "|"
    assert detect_delimiter("a\tb\n1\t2\n") == "\t"
    assert detect_delimiter('a,b\n"x;y",1\n"z;w",2\n') == ","
    assert detect_delimiter("") == ","
    assert detect_decimal("a;b\n1,5;2,25\n3,5;4\n", ";") == ","
    assert detect_decimal("a;b\n1.5;2.25\n", ";") == "."
    assert detect_decimal("a,b\n1,5\n", ",") == "."
    assert detect_decimal("d;v\n01.02.2025;1,5\n", ";") == ","


def test_detect_header_row() -> None:
    assert detect_header_row([["Title"], [], ["a", "b", "c"], ["1", "2", "3"]]) == 2
    assert detect_header_row([["1", "2"], ["3", "4"]]) == 0
    assert detect_header_row([["a", "b"], ["1", "2"]]) == 0
    assert detect_header_row([["T", None, None], [None, None, None], ["x", "y", "z"], [1, 2.5, "q"]]) == 2
    assert detect_header_row([]) == 0


def test_current_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    assert current_limits() == SERVER_LIMITS
    monkeypatch.setenv("DATA_EXPLORER_MAX_MB", "10")
    monkeypatch.setenv("DATA_EXPLORER_MAX_ROWS", "100")
    assert current_limits() == dio.Limits(10 * MiB, 100, 500)
    monkeypatch.setenv("DATA_EXPLORER_MAX_MB", "500")
    monkeypatch.setenv("DATA_EXPLORER_MAX_COLS", "900")
    assert current_limits() == dio.Limits(50 * MiB, 100, 500)
    for bad in ("abc", "0", "-5", "1.5"):
        monkeypatch.setenv("DATA_EXPLORER_MAX_COLS", bad)
        with pytest.raises(ValueError, match="DATA_EXPLORER_MAX_COLS"):
            current_limits()
    monkeypatch.setenv("DATA_EXPLORER_MAX_COLS", "")
    assert current_limits().max_cols == 500
    monkeypatch.setattr(sys, "platform", "emscripten")
    assert current_limits() == BROWSER_LIMITS


def test_samples_registry_and_ids() -> None:
    assert list(SAMPLES) == ["retail_orders", "sensor_readings", "client_survey"]
    assert [info.title for info in SAMPLES.values()] == ["Retail orders", "Sensor readings", "Client survey"]
    with pytest.raises(KeyError):
        load_sample("nope")
    data = sample_bytes("retail_orders")
    first = dataset_id(data, LoadOptions())
    assert len(first) == 16 and all(c in "0123456789abcdef" for c in first)
    assert dataset_id(data, LoadOptions(delimiter=";")) != first


def test_report_is_json_safe(fixture_bytes: Read) -> None:
    _, report = load_table(fixture_bytes("duplicate_and_blank_headers.csv"), "h.csv")
    payload = report.to_dict()
    assert json.loads(json.dumps(payload)) == payload
    assert payload["renamed_columns"][0] == ["", "column_3"]
    assert payload["sha256"] == report.sha256 and len(report.sha256) == 64
