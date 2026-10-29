"""Read CSV and XLSX bytes into a DataFrame: detection, header cleanup, missing values, dates, limits."""

from __future__ import annotations

import codecs
import csv
import dataclasses
import datetime as dt
import hashlib
import io
import itertools
import json
import logging
import math
import os
import re
import sys
import time
import zipfile
from collections import Counter
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from importlib import resources
from pathlib import PurePath
from types import MappingProxyType
from typing import Any, Literal

import pandas as pd
from openpyxl import load_workbook

from data_explorer.core import JSONValue, json_object

logger = logging.getLogger(__name__)

MiB = 1024 * 1024
DETECT_BYTES = MiB
SNIFF_CHARS = 64 * 1024
HEADER_SCAN_ROWS = 20
DECIMAL_SCAN_ROWS = 200
SHEET_SCAN_ROWS = 200
MAX_HEADER_ROW = 19
XLSX_MAX_ROWS = 1_048_576  # the row limit of Excel itself; a <row r=...> beyond it is not a real sheet
XLSX_EXPANSION = 10  # a workbook may declare at most this many times the upload limit once unpacked
XLSX_SHARED_STRINGS = "xl/sharedStrings.xml"  # parsed whole on open; it may declare at most the upload limit
DATE_SHARE_PERCENT = 95
DATE_LIKE_PERCENT = 50
ENCODINGS = ("utf-8", "utf-8-sig", "cp1252", "latin-1")
DELIMITERS = (",", ";", "\t", "|")
DECIMALS = (".", ",")
CSV_SUFFIXES = (".csv", ".tsv", ".txt")
XLSX_SUFFIXES = (".xlsx",)
# The pandas default NA tokens, matched after stripping the cell.
NA_TOKENS = frozenset(
    {
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
    }
)

FileType = Literal["csv", "xlsx"]
LoadErrorCode = Literal[
    "empty_file",
    "too_large",
    "unsupported_type",
    "decode_failed",
    "parse_failed",
    "sheet_not_found",
    "no_columns",
    "too_many_rows",
    "too_many_columns",
]

NUMBER_LIKE = re.compile(r"-?\d+(?:[.,]\d+)+")
DOTTED_DATE = re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}")
COMMA_DECIMAL = re.compile(r"-?\d{1,3}(?:\.\d{3})*,\d+|-?\d+,\d+")
LOOKS_NUMERIC = re.compile(r"[-+]?(?:\d+(?:[.,]\d+)*|\d*[.,]\d+)(?:[eE][-+]?\d+)?")
INT_POINT = re.compile(r"-?\d+")
FLOAT_POINT = re.compile(r"-?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][-+]?\d+)?")
INT_COMMA = re.compile(r"-?[1-9]\d{0,2}(?:\.\d{3})+|-?\d+")  # the first group of thousands starts with 1-9
FLOAT_COMMA = re.compile(r"-?(?:[1-9]\d{0,2}(?:\.\d{3})+|\d+),\d+")
LEADING_ZERO = re.compile(r"-?0\d")
MAX_INT_DIGITS = 15
DATE_HINT = re.compile(r"\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
UTC_OFFSET = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})")
DATE_LIKE = re.compile(r"\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{4}")
# Tried in order: ISO forms first, then day-first forms. There is no month-first format on purpose.
DATE_FORMATS = (
    ("%Y-%m-%d", re.compile(r"\d{4}-\d{2}-\d{2}")),
    ("%Y-%m-%d %H:%M", re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}")),
    ("%Y-%m-%d %H:%M:%S", re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}")),
    ("%Y-%m-%d %H:%M:%S.%f", re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\.\d{1,6}")),
    ("%Y-%m-%dT%H:%M", re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")),
    ("%Y-%m-%dT%H:%M:%S", re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}")),
    ("%Y-%m-%dT%H:%M:%S.%f", re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\.\d{1,6}")),
    ("%d/%m/%Y", re.compile(r"\d{1,2}/\d{1,2}/\d{4}")),
    ("%d/%m/%Y %H:%M", re.compile(r"\d{1,2}/\d{1,2}/\d{4} \d{1,2}:\d{2}")),
    ("%d/%m/%Y %H:%M:%S", re.compile(r"\d{1,2}/\d{1,2}/\d{4} \d{1,2}:\d{2}:\d{2}")),
    ("%d.%m.%Y", re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}")),
    ("%d.%m.%Y %H:%M", re.compile(r"\d{1,2}\.\d{1,2}\.\d{4} \d{1,2}:\d{2}")),
    ("%d.%m.%Y %H:%M:%S", re.compile(r"\d{1,2}\.\d{1,2}\.\d{4} \d{1,2}:\d{2}:\d{2}")),
)


@dataclass(frozen=True)
class Limits:
    max_bytes: int
    max_rows: int
    max_cols: int


SERVER_LIMITS = Limits(50 * MiB, 1_000_000, 500)
BROWSER_LIMITS = Limits(20 * MiB, 200_000, 200)


def _env_positive_int(name: str) -> int | None:
    raw = os.environ.get(name)
    if raw is None or raw.strip() == "":
        return None
    if not re.fullmatch(r"[1-9][0-9]*", raw.strip()):
        raise ValueError(f"{name} must be a positive integer, got {raw!r}.")
    return int(raw.strip())


def current_limits() -> Limits:
    """Browser limits under Pyodide; otherwise the server limits, lowered by environment variables."""
    if sys.platform == "emscripten":
        return BROWSER_LIMITS
    megabytes = _env_positive_int("DATA_EXPLORER_MAX_MB")
    rows = _env_positive_int("DATA_EXPLORER_MAX_ROWS")
    cols = _env_positive_int("DATA_EXPLORER_MAX_COLS")
    return Limits(
        max_bytes=min(megabytes * MiB, SERVER_LIMITS.max_bytes) if megabytes else SERVER_LIMITS.max_bytes,
        max_rows=min(rows, SERVER_LIMITS.max_rows) if rows else SERVER_LIMITS.max_rows,
        max_cols=min(cols, SERVER_LIMITS.max_cols) if cols else SERVER_LIMITS.max_cols,
    )


@dataclass(frozen=True)
class LoadOptions:
    """Reading options; ``None`` means detect. ``header_row`` is 0-based."""

    encoding: str | None = None
    delimiter: str | None = None
    decimal: str | None = None
    sheet: str | None = None
    header_row: int | None = None

    def __post_init__(self) -> None:
        if self.encoding is not None and self.encoding not in ENCODINGS:
            raise ValueError(f"Unsupported encoding {self.encoding!r}; use one of {', '.join(ENCODINGS)}.")
        if self.delimiter is not None and self.delimiter not in DELIMITERS:
            raise ValueError(f"Unsupported delimiter {self.delimiter!r}.")
        if self.decimal is not None and self.decimal not in DECIMALS:
            raise ValueError(f"Unsupported decimal separator {self.decimal!r}.")
        if self.header_row is not None and not 0 <= self.header_row <= MAX_HEADER_ROW:
            raise ValueError(f"header_row must be between 0 and {MAX_HEADER_ROW}.")

    def to_json(self) -> str:
        return json.dumps(dataclasses.asdict(self), sort_keys=True, separators=(",", ":"))

    @classmethod
    def from_json(cls, text: str) -> LoadOptions:
        data = json.loads(text)
        if not isinstance(data, dict):
            raise ValueError("LoadOptions JSON must be an object.")
        return cls(
            encoding=data.get("encoding"),
            delimiter=data.get("delimiter"),
            decimal=data.get("decimal"),
            sheet=data.get("sheet"),
            header_row=data.get("header_row"),
        )


DEFAULT_OPTIONS = LoadOptions()


@dataclass(frozen=True)
class LoadReport:
    file_name: str
    file_type: FileType
    size_bytes: int
    sha256: str
    encoding: str | None
    delimiter: str | None
    decimal: str | None
    sheet: str | None
    sheets: tuple[str, ...]
    header_row: int
    n_rows: int
    n_cols: int
    renamed_columns: tuple[tuple[str, str], ...]
    date_columns: tuple[str, ...]
    blank_to_missing: int
    warnings: tuple[str, ...]

    def to_dict(self) -> dict[str, JSONValue]:
        return json_object(self)


class LoadError(ValueError):
    """A file that cannot be loaded; ``message`` is shown to the user as it is."""

    def __init__(self, code: LoadErrorCode, message: str) -> None:
        super().__init__(message)
        self.code: LoadErrorCode = code
        self.message = message


@dataclass(frozen=True)
class SampleInfo:
    key: str
    title: str
    description: str
    file_name: str
    defects: tuple[str, ...]


SAMPLES: Mapping[str, SampleInfo] = MappingProxyType(
    {
        "retail_orders": SampleInfo(
            key="retail_orders",
            title="Retail orders",
            description=(
                "Orders of a fictional homeware chain in eight Lombard cities, January 2024 to June 2025. "
                "UTF-8 with BOM, semicolons, decimal commas, dates as dd/mm/yyyy."
            ),
            file_name="retail_orders.csv",
            defects=(
                "60 exact duplicate rows",
                "180 missing discount_pct values",
                "15 units values between 120 and 400 (IQR outliers)",
                "country is always IT",
            ),
        ),
        "sensor_readings": SampleInfo(
            key="sensor_readings",
            title="Sensor readings",
            description=(
                "Hourly climate readings of six office sensors over 60 days from 1 March 2025. "
                "cp1252, commas, decimal points, ISO timestamps."
            ),
            file_name="sensor_readings.csv",
            defects=(
                "sensor MX-4 offline for 48 hours (temperature, humidity and CO2 missing)",
                "30 further missing humidity_pct values",
                "12 co2_ppm sentinel values of 9999",
                "firmware is always 2.4.1",
            ),
        ),
        "client_survey": SampleInfo(
            key="client_survey",
            title="Client survey",
            description=(
                "Satisfaction survey of a fictional service provider: an Excel workbook with a Responses "
                "and a Notes sheet; the header is on row 4, below two title rows and a blank row."
            ),
            file_name="client_survey.xlsx",
            defects=(
                "24 exact duplicate rows",
                "37 missing satisfaction values",
                "840 empty comments",
                "10 response_time_s values above 20,000",
                "survey_wave is always 2025",
            ),
        ),
    }
)


def sample_bytes(key: str) -> bytes:
    if key not in SAMPLES:
        raise KeyError(f"Unknown sample {key!r}")
    return resources.files("data_explorer").joinpath("samples").joinpath(SAMPLES[key].file_name).read_bytes()


def load_sample(key: str, limits: Limits | None = None) -> tuple[pd.DataFrame, LoadReport]:
    data = sample_bytes(key)
    return load_table(data, SAMPLES[key].file_name, DEFAULT_OPTIONS, limits)


def dataset_id(data: bytes, options: LoadOptions) -> str:
    return hashlib.sha256(data + options.to_json().encode("utf-8")).hexdigest()[:16]


# --- detection -------------------------------------------------------------------------------


def detect_encoding(data: bytes) -> str:
    """BOM: utf-8-sig; strict UTF-8 on the first MiB: utf-8; strict cp1252: cp1252; else latin-1."""
    if data.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    head = data[:DETECT_BYTES]
    try:
        codecs.getincrementaldecoder("utf-8")(errors="strict").decode(head, final=False)
        return "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        head.decode("cp1252")
        return "cp1252"
    except UnicodeDecodeError:
        return "latin-1"


def detect_delimiter(sample: str) -> str:
    """csv.Sniffer restricted to , ; tab | on the first 64 KiB; then a line-count vote; else a comma."""
    head = sample[:SNIFF_CHARS]
    if len(sample) > SNIFF_CHARS and "\n" in head:
        head = head[: head.rfind("\n")]
    lines = [line for line in head.splitlines() if line.strip()]
    if not lines:
        return ","
    sniffer = csv.Sniffer()
    sniffer.preferred = ["\t", ";", "|", ","]
    try:
        delimiter = sniffer.sniff("\n".join(lines), delimiters="".join(DELIMITERS)).delimiter
        if delimiter in DELIMITERS:
            return delimiter
    except csv.Error:
        pass
    best, best_lines = ",", 0
    window = lines[:HEADER_SCAN_ROWS]
    for candidate in DELIMITERS:
        counts = [line.count(candidate) for line in window if line.count(candidate) > 0]
        if not counts:
            continue
        lines_with_mode = Counter(counts).most_common(1)[0][1]
        if lines_with_mode > best_lines:
            best, best_lines = candidate, lines_with_mode
    return best


def detect_decimal(sample: str, delimiter: str) -> str:
    """',' when the delimiter is not ',' and at least 60 % of number-like fields use a decimal comma."""
    if delimiter == ",":
        return "."
    number_like = comma = 0
    reader = csv.reader(io.StringIO(sample), delimiter=delimiter)
    for row in itertools.islice(reader, DECIMAL_SCAN_ROWS):
        for field in row:
            value = field.strip()
            if not NUMBER_LIKE.fullmatch(value) or DOTTED_DATE.fullmatch(value):
                continue
            number_like += 1
            if COMMA_DECIMAL.fullmatch(value):
                comma += 1
    return "," if number_like and comma * 100 >= 60 * number_like else "."


def _is_blank(cell: object) -> bool:
    return cell is None or (isinstance(cell, str) and cell.strip() == "")


def _is_text_cell(cell: object) -> bool:
    return isinstance(cell, str) and not LOOKS_NUMERIC.fullmatch(cell.strip())


def detect_header_row(rows: Sequence[Sequence[object]]) -> int:
    """First of the first 20 rows whose non-empty cells are >= 50 % of the widest row and all text."""
    window = [list(row) for row in rows[:HEADER_SCAN_ROWS]]
    widest = max((sum(1 for cell in row if not _is_blank(cell)) for row in window), default=0)
    for index, row in enumerate(window):
        cells = [cell for cell in row if not _is_blank(cell)]
        if cells and 2 * len(cells) >= widest and all(_is_text_cell(cell) for cell in cells):
            return index
    return 0


def _check_unpacked_size(data: bytes, file_name: str, limits: Limits) -> None:
    """Refuse a workbook that declares too much once unpacked.

    Two caps, both on the sizes in the zip directory (``zipfile`` stops a member at its declared size, so
    they bound the work): all parts together at most 10 times the upload limit, and ``xl/sharedStrings.xml``
    at most the upload limit. openpyxl parses the shared strings completely on open, and Deflate reaches
    about 1,000:1, so a few hundred KB of that part cost tens of seconds and over a gigabyte.
    """
    cap = XLSX_EXPANSION * limits.max_bytes
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items: list[zipfile.ZipInfo] = archive.infolist()
    declared: int = sum(item.file_size for item in items)
    if declared > cap:
        raise LoadError(
            "too_large",
            f"{file_name} unpacks to {_size_text(declared)}; the limit is {_size_text(cap)} unpacked "
            f"({XLSX_EXPANSION} times the upload limit).",
        )
    for item in items:
        if item.filename == XLSX_SHARED_STRINGS and item.file_size > limits.max_bytes:
            raise LoadError(
                "too_large",
                f"{file_name} has {XLSX_SHARED_STRINGS} unpacking to {_size_text(item.file_size)}; "
                f"the limit is {_size_text(limits.max_bytes)} unpacked for that part (the upload limit).",
            )


def _open_workbook(data: bytes, file_name: str, limits: Limits | None = None) -> Any:
    active = limits if limits is not None else current_limits()
    try:
        _check_unpacked_size(data, file_name, active)
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except LoadError:
        raise
    except Exception as exc:  # openpyxl leaks zipfile, zlib, XML and index errors from damaged files
        raise LoadError("parse_failed", f"{file_name} could not be read as an Excel workbook.") from exc


def _no_worksheets(file_name: str) -> LoadError:
    return LoadError("no_columns", f"{file_name} has no worksheets, so it contains no columns.")


def _sheet_titles(workbook: Any) -> tuple[str, ...]:
    return tuple(str(sheet.title) for sheet in workbook.worksheets)


def _pick_sheet(workbook: Any, file_name: str = "the workbook") -> str:
    titles = _sheet_titles(workbook)
    if not titles:
        raise _no_worksheets(file_name)
    best, best_count = titles[0], -1
    for sheet in workbook.worksheets:
        sheet.reset_dimensions()  # a wrong <dimension> must not pad the rows
        count = 0
        for row in sheet.iter_rows(max_row=SHEET_SCAN_ROWS, values_only=True):
            count += sum(1 for cell in row if not _is_blank(cell))
        if count > best_count:
            best, best_count = str(sheet.title), count
    return best


def list_sheets(
    data: bytes, file_name: str = "the workbook", limits: Limits | None = None
) -> tuple[str, ...]:
    workbook = _open_workbook(data, file_name, limits)
    try:
        return _sheet_titles(workbook)
    finally:
        workbook.close()


def detect_sheet(data: bytes, limits: Limits | None = None) -> str:
    """The sheet with the most non-empty cells in its first 200 rows; ties go to the first sheet."""
    workbook = _open_workbook(data, "the workbook", limits)
    try:
        return _pick_sheet(workbook)
    except LoadError:
        raise
    except Exception as exc:
        raise LoadError("parse_failed", f"The workbook could not be read: {_first_line(exc)}") from exc
    finally:
        workbook.close()


# --- column inference ------------------------------------------------------------------------


@dataclass
class _Column:
    series: pd.Series[Any]
    blank: int
    warnings: list[str]


def _clean_text_cells(raw: Sequence[object]) -> tuple[list[str | None], int]:
    """Missing: None, NaN, blank-after-strip (counted) and the NA tokens; text is kept as written."""
    cleaned: list[str | None] = []
    blank = 0
    for value in raw:
        if value is None or (isinstance(value, float) and math.isnan(value)):
            cleaned.append(None)
            continue
        text = str(value)
        stripped = text.strip()
        if stripped == "":
            if text != "":
                blank += 1
            cleaned.append(None)
        elif stripped in NA_TOKENS:
            cleaned.append(None)
        else:
            cleaned.append(text)
    return cleaned, blank


def _parse_numbers(values: Sequence[str], decimal: str) -> tuple[list[int | float], bool] | None:
    numbers: list[int | float] = []
    all_int = True
    for value in values:
        if LEADING_ZERO.match(value):
            return None
        if decimal == ",":
            if INT_COMMA.fullmatch(value):
                numbers.append(int(value.replace(".", "")))
            elif FLOAT_COMMA.fullmatch(value):
                numbers.append(float(value.replace(".", "").replace(",", ".")))
                all_int = False
            else:
                return None
        elif INT_POINT.fullmatch(value):
            numbers.append(int(value))
        elif FLOAT_POINT.fullmatch(value):
            numbers.append(float(value))
            all_int = False
        else:
            return None
        if all_int and len(value.lstrip("-").replace(".", "")) > MAX_INT_DIGITS:
            return None
    return numbers, all_int


def _parse_dates(
    stripped: Sequence[str | None], total: int, name: str
) -> tuple[pd.Series[Any] | None, list[str]]:
    """The parsed column and its warnings; no series when the column stays text (maybe with a warning)."""
    present = [value for value in stripped if value is not None]
    if not any(DATE_HINT.match(value) for value in present[:50]):
        return None, []
    values = pd.Series(list(stripped), dtype="object")
    if sum(1 for value in present if UTC_OFFSET.fullmatch(value)) * 100 >= DATE_SHARE_PERCENT * total:
        parsed = pd.to_datetime(values, format="ISO8601", utc=True, errors="coerce").dt.tz_localize(None)
        warnings = [f"Column '{name}': times with a UTC offset were converted to UTC."]
        failed = total - int(parsed.notna().sum())
        if failed:
            warnings.append(
                f"Column '{name}': {failed} of {total} values could not be read as dates and are now missing."
            )
        return parsed.astype("datetime64[us]"), warnings
    for fmt, pattern in DATE_FORMATS:
        if not any(pattern.fullmatch(value) for value in present):
            continue
        parsed = pd.to_datetime(values, format=fmt, errors="coerce")
        valid = int(parsed.notna().sum())
        if valid * 100 >= DATE_SHARE_PERCENT * total:
            warnings = []
            if valid < total:
                warnings.append(
                    f"Column '{name}': {total - valid} of {total} values could not be read as dates ({fmt}) "
                    "and are now missing."
                )
            return parsed.astype("datetime64[us]"), warnings
    if sum(1 for value in present if DATE_LIKE.match(value)) * 100 >= DATE_LIKE_PERCENT * total:
        return None, [f"Column '{name}' looks like dates but no known format matched; kept as text."]
    return None, []


def _infer_strings(
    cleaned: Sequence[str | None], decimal: str, name: str
) -> tuple[pd.Series[Any], list[str]]:
    """Booleans, then numbers, then dates; otherwise text kept as written. All-missing stays text."""
    stripped = [None if value is None else value.strip() for value in cleaned]
    present = [value for value in stripped if value is not None]
    if not present:
        return pd.Series(list(cleaned), dtype="str"), []
    if {value.casefold() for value in present} <= {"true", "false"}:
        flags = [None if value is None else value.casefold() == "true" for value in stripped]
        return pd.Series(flags, dtype="boolean" if None in flags else "bool"), []
    parsed = _parse_numbers(present, decimal)
    if parsed is not None:
        numbers, all_int = parsed
        values: list[int | float | None] = []
        iterator = iter(numbers)
        for value in stripped:
            values.append(None if value is None else next(iterator))
        missing = len(present) < len(stripped)
        if all_int:
            return pd.Series(values, dtype="Int64" if missing else "int64"), []
        return pd.Series(values, dtype="float64"), []
    dates, date_warnings = _parse_dates(stripped, len(present), name)
    if dates is not None:
        return dates, date_warnings
    return pd.Series(list(cleaned), dtype="str"), date_warnings


def _infer_text_column(raw: Sequence[object], decimal: str, name: str) -> _Column:
    cleaned, blank = _clean_text_cells(raw)
    series, warnings = _infer_strings(cleaned, decimal, name)
    return _Column(series, blank, warnings)


def _cell_text(value: object) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, dt.datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return str(value)


def _infer_cell_column(raw: Sequence[object], name: str) -> _Column:
    """Typed Excel cells keep their type when a column has one type; mixed columns go through text."""
    cleaned: list[object] = []
    blank = 0
    for value in raw:
        if isinstance(value, str):
            stripped = value.strip()
            if stripped == "":
                blank += 1
                cleaned.append(None)
                continue
            if stripped in NA_TOKENS:
                cleaned.append(None)
                continue
        cleaned.append(value)
    present = [value for value in cleaned if value is not None]
    missing = len(present) < len(cleaned)
    if present and all(isinstance(value, bool) for value in present):
        return _Column(pd.Series(cleaned, dtype="boolean" if missing else "bool"), blank, [])
    if present and all(isinstance(value, int) and not isinstance(value, bool) for value in present):
        return _Column(pd.Series(cleaned, dtype="Int64" if missing else "int64"), blank, [])
    if present and all(isinstance(value, int | float) and not isinstance(value, bool) for value in present):
        floats = [None if value is None else float(value) for value in cleaned]  # type: ignore[arg-type]
        return _Column(pd.Series(floats, dtype="float64"), blank, [])
    if present and all(isinstance(value, dt.datetime | dt.date) for value in present):
        stamps = pd.to_datetime(pd.Series(cleaned, dtype="object"), errors="coerce")
        return _Column(stamps.astype("datetime64[us]"), blank, [])
    texts = [None if value is None else _cell_text(value) for value in cleaned]
    series, warnings = _infer_strings(texts, ".", name)
    return _Column(series, blank, warnings)


# --- loading ---------------------------------------------------------------------------------


@dataclass
class _Table:
    raw_names: list[object]
    columns: list[_Column]
    n_rows: int
    header_row: int
    encoding: str | None = None
    delimiter: str | None = None
    decimal: str | None = None
    sheet: str | None = None
    sheets: tuple[str, ...] = ()


def _size_text(size: int) -> str:
    if size % MiB == 0:
        return f"{size // MiB} MB"
    if size >= 100 * 1024:
        return f"{size / MiB:.1f} MB"
    return f"{size:,} bytes"


def _too_many_columns(file_name: str, n_cols: int, limits: Limits, *, at_least: bool = False) -> LoadError:
    found = f"at least {n_cols:,}" if at_least else f"{n_cols:,}"
    return LoadError(
        "too_many_columns", f"{file_name} has {found} columns; the limit is {limits.max_cols:,} columns."
    )


def _check_shape(file_name: str, n_rows: int, n_cols: int, limits: Limits) -> None:
    """The row count may be a lower bound: both readers stop one row past the limit."""
    if n_rows > limits.max_rows:
        raise LoadError(
            "too_many_rows",
            f"{file_name} has more than {limits.max_rows:,} rows; the limit is {limits.max_rows:,} rows.",
        )
    if n_cols > limits.max_cols:
        raise _too_many_columns(file_name, n_cols, limits)


def _first_line(exc: Exception) -> str:
    text = str(exc).strip().splitlines()
    return (text[0] if text else type(exc).__name__)[:200]


def _load_csv(data: bytes, file_name: str, options: LoadOptions, limits: Limits) -> _Table:
    encoding = options.encoding or detect_encoding(data)
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError as exc:
        raise LoadError(
            "decode_failed",
            f"{file_name} could not be decoded as {encoding} (byte {exc.start + 1:,}). "
            "Choose another encoding under Reading options.",
        ) from exc
    if "\x00" in text:
        raise LoadError(
            "decode_failed", f"{file_name} does not look like a text file (it contains NUL bytes)."
        )
    text = text.removeprefix(chr(0xFEFF))
    if not text.strip():
        raise LoadError("no_columns", f"{file_name} contains no columns.")
    sample = text[:DETECT_BYTES]
    delimiter = options.delimiter or detect_delimiter(sample)
    decimal = options.decimal or detect_decimal(sample, delimiter)
    try:
        head = list(itertools.islice(csv.reader(io.StringIO(sample), delimiter=delimiter), HEADER_SCAN_ROWS))
    except csv.Error as exc:
        raise LoadError(
            "parse_failed", f"{file_name} could not be read as a table: {_first_line(exc)}"
        ) from exc
    header_row = options.header_row if options.header_row is not None else detect_header_row(head)
    if header_row < len(head) and len(head[header_row]) > limits.max_cols:
        # Refused before the parse: a header of a million fields would cost gigabytes. The sample ends at
        # DETECT_BYTES, so a header line cut there is only known to have at least this many fields.
        cut = len(sample) < len(text) and header_row == len(head) - 1
        raise _too_many_columns(file_name, len(head[header_row]), limits, at_least=cut)
    try:
        frame = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            header=None,
            dtype=object,
            keep_default_na=False,
            na_filter=False,
            skiprows=header_row,
            skip_blank_lines=True,
            nrows=limits.max_rows + 2,  # the header and one row past the limit; the parse stops there
            engine="c",
        )
    except pd.errors.EmptyDataError as exc:
        raise LoadError("no_columns", f"{file_name} contains no columns.") from exc
    except (pd.errors.ParserError, ValueError, csv.Error) as exc:
        raise LoadError(
            "parse_failed", f"{file_name} could not be read as a table: {_first_line(exc)}"
        ) from exc
    n_rows, n_cols = len(frame) - 1, int(frame.shape[1])
    _check_shape(file_name, n_rows, n_cols, limits)
    raw_names = list(frame.iloc[0].tolist())
    columns = []
    for position in range(n_cols):
        name = "" if raw_names[position] is None else str(raw_names[position]).strip()
        columns.append(
            _infer_text_column(frame.iloc[1:, position].tolist(), decimal, name or f"column_{position + 1}")
        )
    return _Table(
        raw_names, columns, n_rows, header_row, encoding=encoding, delimiter=delimiter, decimal=decimal
    )


def _last_filled(row: Sequence[object]) -> int:
    for index in range(len(row) - 1, -1, -1):
        if not _is_blank(row[index]):
            return index + 1
    return 0


def _bounded_rows(worksheet: Any, max_cols: int) -> Iterator[tuple[list[object], int]]:
    """Rows of a sheet read lazily, each cut to ``max_cols`` cells, with the width its cells need.

    The width is 0 for a row that fits; for a longer row it is the position of its last filled cell, which
    can exceed ``max_cols`` (too many columns) or not (blank padding). Row widths are what the cells in the
    XML say: the sheet's ``<dimension>`` is dropped first, because a wrong one pads every row to that width.
    """
    worksheet.reset_dimensions()
    for row in worksheet.iter_rows(max_row=XLSX_MAX_ROWS, values_only=True):
        if len(row) > max_cols:
            yield list(row[:max_cols]), _last_filled(row)
        else:
            yield list(row), 0


def _read_xlsx_rows(
    worksheet: Any, file_name: str, options: LoadOptions, limits: Limits
) -> tuple[int, list[object], list[list[object]], int]:
    """Header row, header cells, non-blank body rows and the widest row; stops once the row limit is passed."""
    stream = _bounded_rows(worksheet, limits.max_cols)
    head = list(itertools.islice(stream, HEADER_SCAN_ROWS))
    header_row = (
        options.header_row if options.header_row is not None else detect_header_row([r for r, _ in head])
    )
    if header_row >= len(head):
        raise LoadError("no_columns", f"{file_name} contains no columns.")
    names, wide = head[header_row]
    body: list[list[object]] = []
    for row, row_wide in itertools.chain(head[header_row + 1 :], stream):
        wide = max(wide, row_wide)
        if _last_filled(row) > 0:
            body.append(row)
            if len(body) > limits.max_rows:
                raise LoadError(
                    "too_many_rows",
                    f"{file_name} has more than {limits.max_rows:,} rows; the limit is {limits.max_rows:,} rows.",
                )
    return header_row, names, body, wide


def _load_xlsx(data: bytes, file_name: str, options: LoadOptions, limits: Limits) -> _Table:
    workbook = _open_workbook(data, file_name, limits)
    try:
        sheets = _sheet_titles(workbook)
        if not sheets:
            raise _no_worksheets(file_name)
        if options.sheet is not None and options.sheet not in sheets:
            raise LoadError(
                "sheet_not_found",
                f"{file_name} has no sheet named {options.sheet!r}; it has {', '.join(sheets)}.",
            )
        sheet = options.sheet if options.sheet is not None else _pick_sheet(workbook, file_name)
        header_row, names, body, wide = _read_xlsx_rows(workbook[sheet], file_name, options, limits)
    except LoadError:
        raise
    except Exception as exc:  # openpyxl leaks zipfile, zlib, XML and index errors from damaged files
        raise LoadError(
            "parse_failed", f"{file_name} could not be read as an Excel workbook: {_first_line(exc)}"
        ) from exc
    finally:
        workbook.close()
    width = max([_last_filled(names), wide] + [_last_filled(row) for row in body])
    if width == 0:
        raise LoadError("no_columns", f"{file_name} contains no columns.")
    _check_shape(file_name, len(body), width, limits)
    raw_names: list[object] = [names[j] if j < len(names) else None for j in range(width)]
    columns = []
    for position in range(width):
        values = [row[position] if position < len(row) else None for row in body]
        label = "" if raw_names[position] is None else str(raw_names[position]).strip()
        columns.append(_infer_cell_column(values, label or f"column_{position + 1}"))
    return _Table(raw_names, columns, len(body), header_row, sheet=sheet, sheets=sheets)


def _clean_headers(raw: Sequence[tuple[int, str]]) -> tuple[list[str], list[tuple[str, str]]]:
    """Strip names, name blanks column_<n> (1-based position), number repeats _2, _3."""
    used: set[str] = set()
    names: list[str] = []
    renamed: list[tuple[str, str]] = []
    for position, original in raw:
        base = original.strip() or f"column_{position + 1}"
        name, suffix = base, 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        names.append(name)
        if name != original:
            renamed.append((original, name))
    return names, renamed


def _assemble(file_name: str, table: _Table) -> tuple[pd.DataFrame, list[tuple[str, str]], list[str]]:
    kept: list[tuple[int, str, _Column]] = []
    warnings: list[str] = []
    for position, (raw, column) in enumerate(zip(table.raw_names, table.columns, strict=True)):
        original = "" if raw is None else str(raw)
        if original.strip() == "" and table.n_rows > 0 and bool(column.series.isna().all()):
            warnings.append(f"Column {position + 1} has no header and no values and was removed.")
            continue
        kept.append((position, original, column))
    if not kept:
        raise LoadError("no_columns", f"{file_name} contains no columns.")
    names, renamed = _clean_headers([(position, original) for position, original, _ in kept])
    frame = pd.DataFrame({name: column.series for name, (_, _, column) in zip(names, kept, strict=True)})
    frame.index = pd.RangeIndex(len(frame))
    for _, _, column in kept:
        warnings.extend(column.warnings)
    return frame, renamed, warnings


def load_table(
    data: bytes, file_name: str, options: LoadOptions = DEFAULT_OPTIONS, limits: Limits | None = None
) -> tuple[pd.DataFrame, LoadReport]:
    """Load CSV or XLSX bytes. Checks run in the order empty, size, type, decode, parse, shape."""
    started = time.perf_counter()
    active = limits if limits is not None else current_limits()
    if len(data) == 0:
        raise LoadError("empty_file", f"{file_name} is empty.")
    if len(data) > active.max_bytes:
        raise LoadError(
            "too_large",
            f"{file_name} is {_size_text(len(data))}; the limit is {_size_text(active.max_bytes)}.",
        )
    suffix = PurePath(file_name).suffix.lower()
    file_type: FileType
    if suffix in CSV_SUFFIXES:
        file_type = "csv"
        table = _load_csv(data, file_name, options, active)
    elif suffix in XLSX_SUFFIXES:
        file_type = "xlsx"
        table = _load_xlsx(data, file_name, options, active)
    else:
        raise LoadError(
            "unsupported_type",
            f"{file_name}: this file type is not supported. Use .csv, .tsv, .txt or .xlsx; "
            "save legacy .xls workbooks as .xlsx first.",
        )
    frame, renamed, warnings = _assemble(file_name, table)
    report = LoadReport(
        file_name=file_name,
        file_type=file_type,
        size_bytes=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        encoding=table.encoding,
        delimiter=table.delimiter,
        decimal=table.decimal,
        sheet=table.sheet,
        sheets=table.sheets,
        header_row=table.header_row,
        n_rows=len(frame),
        n_cols=len(frame.columns),
        renamed_columns=tuple(renamed),
        date_columns=tuple(
            str(name) for name in frame.columns if pd.api.types.is_datetime64_dtype(frame[name])
        ),
        blank_to_missing=sum(column.blank for column in table.columns),
        warnings=tuple(warnings),
    )
    logger.info(
        "load file_type=%s size_bytes=%d rows=%d cols=%d duration_ms=%.0f",
        file_type,
        len(data),
        report.n_rows,
        report.n_cols,
        (time.perf_counter() - started) * 1000,
    )
    return frame, report
