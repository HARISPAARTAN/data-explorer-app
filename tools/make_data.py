"""Generate the three sample datasets and the test fixtures (seeded and deterministic).

    uv run python tools/make_data.py                write src/data_explorer/samples/
    uv run python tools/make_data.py --fixtures     also write fixtures/
    uv run python tools/make_data.py --samples-dir DIR --fixtures --fixtures-dir DIR

Seed 20261029 with numpy's default_rng (one child stream per dataset). Defects are planted in exact
numbers on distinct rows; duplicates copy rows without any other defect, so every count in
docs/data-model.md is independent of the others. Cities only: no person names, no company names,
no addresses.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import math
import re
import zipfile
from collections.abc import Mapping, Sequence
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

SEED = 20261029
ROOT = Path(__file__).resolve().parents[1]
SAMPLES_DIR = ROOT / "src" / "data_explorer" / "samples"
FIXTURES_DIR = ROOT / "fixtures"
FIXED_TIME = dt.datetime(2026, 10, 29, 0, 0, 0)
ZIP_TIME = (2026, 10, 29, 0, 0, 0)
MODIFIED_XML = re.compile(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")

STORES = ("Milan", "Bergamo", "Brescia", "Como", "Monza", "Pavia", "Varese", "Lecco")
STORE_WEIGHTS = (0.22, 0.14, 0.13, 0.11, 0.12, 0.09, 0.10, 0.09)
CATEGORIES = ("Kitchen", "Textiles", "Lighting", "Storage", "Decor", "Garden")
PRICE_RANGES = {
    "Kitchen": (4, 60),
    "Textiles": (6, 80),
    "Lighting": (12, 150),
    "Storage": (5, 90),
    "Decor": (3, 70),
    "Garden": (4, 120),
}
DISCOUNTS = (0, 5, 10, 15, 20)
DISCOUNT_WEIGHTS = (0.5, 0.2, 0.15, 0.1, 0.05)
PAYMENTS = ("Card", "Cash", "Voucher")
PAYMENT_WEIGHTS = (0.7, 0.22, 0.08)
CUSTOMER_TYPES = ("New", "Returning")
CUSTOMER_WEIGHTS = (0.35, 0.65)
RETAIL_COLUMNS = (
    "order_id",
    "order_date",
    "store",
    "category",
    "units",
    "unit_price_eur",
    "discount_pct",
    "revenue_eur",
    "payment",
    "customer_type",
    "country",
)

ROOMS = {
    "MX-1": "Caffè",
    "MX-2": "Meeting room A",
    "MX-3": "Meeting room B",
    "MX-4": "Open space north",
    "MX-5": "Open space south",
    "MX-6": "Library",
}
CAPACITY = {"MX-1": 12, "MX-2": 10, "MX-3": 8, "MX-4": 30, "MX-5": 30, "MX-6": 15}
OFFLINE_SENSOR = "MX-4"
OFFLINE_START = dt.datetime(2025, 3, 21, 6, 0)
OFFLINE_HOURS = 48
SENSOR_COLUMNS = (
    "timestamp",
    "sensor_id",
    "room",
    "temperature_c",
    "humidity_pct",
    "co2_ppm",
    "occupancy",
    "battery_pct",
    "firmware",
)

INDUSTRIES = ("Manufacturing", "Retail", "Logistics", "Healthcare", "Finance", "Public sector")
COMPANY_SIZES = ("1-9", "10-49", "50-249", "250+")
REGIONS = ("North", "Centre", "South", "Islands")
COMMENTS = (
    "Quick and clear answers.",
    "The portal is easy to use.",
    "Response times could be shorter.",
    "Invoices were correct and on time.",
    "We would like more reporting options.",
    "Support solved the issue on the first call.",
    "Onboarding took longer than planned.",
    "Good value for the price.",
    "The mobile view needs work.",
    "Helpful account team.",
    "Documentation is hard to find.",
    "Smooth renewal process.",
)
QUESTIONS = (
    (
        "satisfaction",
        "Overall, how satisfied are you with the service? (1 = very unsatisfied, 5 = very satisfied)",
    ),
    ("nps", "How likely are you to recommend the service to a colleague? (0 to 10)"),
    ("would_recommend", "Would you recommend the service today? (Yes or No)"),
    ("response_time_s", "Seconds spent on the survey (recorded automatically)"),
    ("comment", "Is there anything you would like to tell us? (optional)"),
    ("company_size", "How many people work at your organisation?"),
)
SURVEY_COLUMNS = (
    "response_id",
    "submitted_at",
    "industry",
    "company_size",
    "region",
    "satisfaction",
    "nps",
    "would_recommend",
    "response_time_s",
    "comment",
    "survey_wave",
)


def _rngs() -> list[np.random.Generator]:
    """Four independent streams: retail, sensor, survey, fixtures."""
    return [np.random.default_rng(child) for child in np.random.SeedSequence(SEED).spawn(4)]


def _csv_bytes(header: Sequence[str], rows: Sequence[Sequence[str]], delimiter: str, encoding: str) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode(encoding)


def _decimal_comma(value: float) -> str:
    """12.9 -> '12,90'; 1234.5 -> '1.234,50' (Italian Excel export)."""
    return f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _pin_modified(core_xml: bytes) -> bytes:
    """Set dcterms:modified in docProps/core.xml to FIXED_TIME (openpyxl stamps the save time there)."""
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
    pinned, count = MODIFIED_XML.subn(rf"\g<1>{stamp}\g<2>", core_xml.decode("utf-8"))
    assert count == 1, "docProps/core.xml has no dcterms:modified element"
    return pinned.encode("utf-8")


def _normalise_zip(data: bytes) -> bytes:
    """Rewrite an XLSX container so two builds are byte-identical.

    Every entry gets the fixed ZIP_TIME, and docProps/core.xml gets dcterms:modified set to
    FIXED_TIME (openpyxl overwrites it with the save time). All other entries keep their bytes.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = _pin_modified(content)
            target.writestr(fixed, content)
    return out.getvalue()


def _workbook_bytes(workbook: Workbook) -> bytes:
    workbook.properties.creator = "data-explorer-app tools/make_data.py"
    workbook.properties.lastModifiedBy = "data-explorer-app tools/make_data.py"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _normalise_zip(buffer.getvalue())


def cell_values(data: bytes) -> dict[str, list[tuple[object, ...]]]:
    """Every sheet's rows as tuples of cell values, for comparisons that ignore the zip container."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return {
            sheet.title: [tuple(row) for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()


def build_retail(rng: np.random.Generator) -> bytes:
    n = 6_000
    start = dt.date(2024, 1, 1)
    days = np.sort(rng.integers(0, 547, size=n))
    stores = rng.choice(len(STORES), size=n, p=STORE_WEIGHTS)
    categories = rng.integers(0, len(CATEGORIES), size=n)
    units = rng.integers(1, 13, size=n).tolist()
    discounts = rng.choice(len(DISCOUNTS), size=n, p=DISCOUNT_WEIGHTS)
    payments = rng.choice(len(PAYMENTS), size=n, p=PAYMENT_WEIGHTS)
    customers = rng.choice(len(CUSTOMER_TYPES), size=n, p=CUSTOMER_WEIGHTS)
    price_draw = rng.random(size=n)
    order = rng.permutation(n)
    missing_discount = set(order[:180].tolist())
    high_units = order[180:195].tolist()
    duplicate_sources = set(order[195:255].tolist())
    for row, value in zip(high_units, rng.integers(120, 401, size=15).tolist(), strict=True):
        units[row] = value
    rows: list[list[str]] = []
    for i in range(n):
        category = CATEGORIES[int(categories[i])]
        low, high = PRICE_RANGES[category]
        price = math.floor(low + float(price_draw[i]) * (high - low)) + 0.9
        discount = DISCOUNTS[int(discounts[i])]
        revenue = round(units[i] * price * (100 - discount) / 100, 2)
        row = [
            f"ORD-{i + 1:06d}",
            (start + dt.timedelta(days=int(days[i]))).strftime("%d/%m/%Y"),
            STORES[int(stores[i])],
            category,
            str(units[i]),
            _decimal_comma(price),
            "" if i in missing_discount else str(discount),
            _decimal_comma(revenue),
            PAYMENTS[int(payments[i])],
            CUSTOMER_TYPES[int(customers[i])],
            "IT",
        ]
        rows.append(row)
        if i in duplicate_sources:
            rows.append(list(row))
    return _csv_bytes(RETAIL_COLUMNS, rows, ";", "utf-8-sig")


def build_sensor(rng: np.random.Generator) -> bytes:
    sensors = tuple(ROOMS)
    start = dt.datetime(2025, 3, 1)
    hours = 60 * 24
    offline_end = OFFLINE_START + dt.timedelta(hours=OFFLINE_HOURS)
    base = rng.uniform(20.5, 22.5, size=len(sensors))
    records: list[dict[str, str]] = []
    offline: list[int] = []
    for hour in range(hours):
        stamp = start + dt.timedelta(hours=hour)
        working = stamp.weekday() < 5 and 8 <= stamp.hour < 18
        daily = math.sin(2 * math.pi * (stamp.hour - 9) / 24)
        for index, sensor in enumerate(sensors):
            capacity = CAPACITY[sensor]
            occupancy = int(rng.binomial(capacity, 0.45)) if working else 0
            noise = rng.normal(0.0, 1.0, size=3)
            temperature = float(base[index]) + 1.2 * daily + 0.06 * occupancy + 0.3 * float(noise[0])
            humidity = min(75.0, max(25.0, 44.0 + 5.0 * daily + 0.2 * occupancy + 1.5 * float(noise[1])))
            co2 = round(
                min(1400.0, max(400.0, 430.0 + 850.0 * occupancy / capacity + 25.0 * float(noise[2])))
            )
            battery = round(100 - 22 * hour / hours - 2 * index)
            if sensor == OFFLINE_SENSOR and OFFLINE_START <= stamp < offline_end:
                offline.append(len(records))
            records.append(
                {
                    "timestamp": stamp.strftime("%Y-%m-%d %H:%M"),
                    "sensor_id": sensor,
                    "room": ROOMS[sensor],
                    "temperature_c": f"{temperature:.1f}",
                    "humidity_pct": f"{humidity:.1f}",
                    "co2_ppm": str(co2),
                    "occupancy": str(occupancy),
                    "battery_pct": str(battery),
                    "firmware": "2.4.1",
                }
            )
    for row in offline:
        for column in ("temperature_c", "humidity_pct", "co2_ppm"):
            records[row][column] = ""
    offline_rows = set(offline)
    eligible = np.array([i for i in range(len(records)) if i not in offline_rows])
    chosen = eligible[rng.permutation(len(eligible))[:42]].tolist()
    for row in chosen[:30]:
        records[row]["humidity_pct"] = ""
    for row in chosen[30:]:
        records[row]["co2_ppm"] = "9999"
    return _csv_bytes(
        SENSOR_COLUMNS, [[record[c] for c in SENSOR_COLUMNS] for record in records], ",", "cp1252"
    )


def build_survey(rng: np.random.Generator) -> bytes:
    n = 1_200
    start = dt.datetime(2025, 5, 5, 8, 0)
    minutes = np.sort(rng.integers(0, 53 * 24 * 60, size=n))
    industries = rng.integers(0, len(INDUSTRIES), size=n)
    sizes = rng.choice(len(COMPANY_SIZES), size=n, p=(0.3, 0.35, 0.22, 0.13))
    regions = rng.choice(len(REGIONS), size=n, p=(0.4, 0.25, 0.25, 0.1))
    satisfaction = rng.choice(5, size=n, p=(0.05, 0.1, 0.25, 0.35, 0.25)) + 1
    nps = np.clip(np.round(satisfaction * 2 + rng.normal(0.0, 1.2, size=n)), 0, 10).astype(int)
    response = rng.integers(60, 1_801, size=n).tolist()
    comments = rng.integers(0, len(COMMENTS), size=n)
    order = rng.permutation(n)
    missing_satisfaction = set(order[:37].tolist())
    empty_comment = set(order[37:877].tolist())
    long_response = order[877:887].tolist()
    duplicate_sources = set(order[887:911].tolist())
    for row, value in zip(long_response, rng.integers(20_001, 90_001, size=10).tolist(), strict=True):
        response[row] = value
    rows: list[list[object]] = []
    for i in range(n):
        row: list[object] = [
            f"R-{i + 1:04d}",
            start + dt.timedelta(minutes=int(minutes[i])),
            INDUSTRIES[int(industries[i])],
            COMPANY_SIZES[int(sizes[i])],
            REGIONS[int(regions[i])],
            None if i in missing_satisfaction else int(satisfaction[i]),
            int(nps[i]),
            "Yes" if nps[i] >= 7 else "No",
            int(response[i]),
            None if i in empty_comment else COMMENTS[int(comments[i])],
            2025,
        ]
        rows.append(row)
        if i in duplicate_sources:
            rows.append(list(row))
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Responses"
    sheet["A1"] = "Client satisfaction survey 2025"
    sheet["A2"] = "Fictional data generated for a demo; no real organisation took part."
    for column, name in enumerate(SURVEY_COLUMNS, start=1):
        sheet.cell(row=4, column=column, value=name)
    for row_number, row in enumerate(rows, start=5):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=row_number, column=column, value=value)
    notes = workbook.create_sheet("Notes")
    notes.append(("question", "text"))
    for question in QUESTIONS:
        notes.append(question)
    return _workbook_bytes(workbook)


def build_samples() -> dict[str, bytes]:
    retail, sensor, survey, _fixtures = _rngs()
    return {
        "retail_orders.csv": build_retail(retail),
        "sensor_readings.csv": build_sensor(sensor),
        "client_survey.xlsx": build_survey(survey),
    }


def _offset_workbook() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    summary["A1"] = "Summary"
    summary["A2"] = "Rows"
    summary["B2"] = 30
    data = workbook.create_sheet("Data")
    data["A1"] = "Monthly figures"
    data["A2"] = "Fictional demo data"
    for column, name in enumerate(("month", "region", "amount", "units"), start=1):
        data.cell(row=4, column=column, value=name)
    for index in range(30):
        values = (
            f"2025-{index % 12 + 1:02d}",
            REGIONS[index % 4],
            round(100 + index * 7.25, 2),
            index % 9 + 1,
        )
        for column, value in enumerate(values, start=1):
            data.cell(row=5 + index, column=column, value=value)
    return _workbook_bytes(workbook)


def build_fixture_files() -> dict[str, bytes]:
    rng = _rngs()[3]
    words = ("Caffè", "Città", "Perché")
    files: dict[str, bytes] = {}
    files["utf8_comma.csv"] = _csv_bytes(
        ("city", "amount", "day", "label"),
        [(STORES[i % 8], f"{10 + i * 1.5:.1f}", f"2025-03-{i + 1:02d}", words[i % 3]) for i in range(20)],
        ",",
        "utf-8",
    )
    files["utf8_bom_semicolon_decimal_comma.csv"] = _csv_bytes(
        ("city", "amount", "day", "units"),
        [
            (STORES[i % 8], _decimal_comma(1234.56 + i * 10.25), f"{i + 1:02d}/03/2025", str(i % 7 + 1))
            for i in range(20)
        ],
        ";",
        "utf-8-sig",
    )
    files["cp1252_semicolon.csv"] = _csv_bytes(
        ("city", "price_eur", "note"),
        [(STORES[i % 8], f"{12.5 + i:.1f}", ("Prezzo in €", "Caffè", "Più tardi")[i % 3]) for i in range(15)],
        ";",
        "cp1252",
    )
    files["latin1_tab.tsv"] = _csv_bytes(
        ("name", "value", "note"),
        [
            (("Kühl", "Grün", "Süß")[i % 3], f"{i * 0.5:.1f}", "ctrl" + chr(0x81) if i == 0 else "ok")
            for i in range(15)
        ],
        "\t",
        "latin-1",
    )
    files["pipe_delimited.txt"] = _csv_bytes(
        ("id", "city", "amount"),
        [(str(i + 1), STORES[i % 8], f"{10.5 + i:.1f}") for i in range(15)],
        "|",
        "utf-8",
    )
    files["header_offset_two_sheets.xlsx"] = _offset_workbook()
    files["duplicate_and_blank_headers.csv"] = (
        "id,name,,name, amount ,name,\n"
        + "".join(f"{i},a{i},x{i},b{i},{i * 10},c{i},\n" for i in range(1, 6))
    ).encode("utf-8")
    dates = [dt.date(2025, 3, 10) + dt.timedelta(days=i) for i in range(20)]
    junk = ("soon", "later", "n/a", "tbd", "2025", "?")
    rows = []
    for i, day in enumerate(dates):
        rows.append(
            (
                (dt.date(2025, 3, 1) + dt.timedelta(days=i)).isoformat(),
                "31/02/2025" if i == 7 else day.strftime("%d/%m/%Y"),
                day.strftime("%d.%m.%Y"),
                day.strftime("%m/%d/%Y"),
                junk[i % len(junk)],
                f"2025-03-{i + 1:02d} {8 + i % 9:02d}:15",
                f"2025-03-{i + 1:02d}T{8 + i % 9:02d}:15:30",
                f"2025-03-{i + 1:02d}T08:00:00+01:00",
            )
        )
    files["dates_mixed.csv"] = _csv_bytes(
        ("iso", "day_first", "dotted", "month_first", "junk", "iso_minutes", "iso_t", "offset"),
        rows,
        ",",
        "utf-8",
    )
    files["whitespace_cells.csv"] = (
        b"name,score,note\nalpha,1,  padded  \n   ,2,ok\nbeta,  ,\t\ngamma,4, \ndelta,5,\n"
    )
    files["formula_cells.csv"] = (
        b"label,formula,amount\na,=SUM(A1:A2),1\nb,+1+1,-2\nc,-minus,3\nd,@mention,-4\ne,plain,5\n"
    )
    noise = rng.normal(0.0, 1.5, size=30)
    files["two_numeric_columns.csv"] = _csv_bytes(
        ("x", "y"), [(str(i), f"{2 * i + float(noise[i]):.2f}") for i in range(30)], ",", "utf-8"
    )
    files["one_numeric_column.csv"] = _csv_bytes(
        ("label", "value"), [(f"item {i + 1}", f"{3.5 + i * 1.25:.2f}") for i in range(20)], ",", "utf-8"
    )
    files["header_only.csv"] = b"a,b,c\n"
    files["empty.csv"] = b""
    random_bytes = rng.bytes(4096)
    assert b"\x00" in random_bytes
    files["not_a_table.bin"] = random_bytes
    files["legacy.xls"] = bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504)
    return files


def write_files(files: Mapping[str, bytes], out_dir: Path) -> list[str]:
    """Write files into out_dir; an XLSX with the same cell values as the existing one is kept."""
    out_dir.mkdir(parents=True, exist_ok=True)
    report: list[str] = []
    for name, data in sorted(files.items()):
        target = out_dir / name
        if (
            name.endswith(".xlsx")
            and target.exists()
            and cell_values(target.read_bytes()) == cell_values(data)
        ):
            report.append(f"kept {target.relative_to(out_dir.parent)} (same cell values)")
            continue
        target.write_bytes(data)
        report.append(f"wrote {target.relative_to(out_dir.parent)} ({len(data):,} bytes)")
    return report


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate the sample datasets and test fixtures.")
    parser.add_argument("--samples-dir", type=Path, default=SAMPLES_DIR)
    parser.add_argument("--fixtures", action="store_true", help="also write the test fixtures")
    parser.add_argument("--fixtures-dir", type=Path, default=FIXTURES_DIR)
    args = parser.parse_args(argv)
    for line in write_files(build_samples(), args.samples_dir):
        print(f"make_data: {line}")
    if args.fixtures:
        for line in write_files(build_fixture_files(), args.fixtures_dir):
            print(f"make_data: {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
